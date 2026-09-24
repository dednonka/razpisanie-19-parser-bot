"""
Телеграм-бот расписания.
Каждые 10 минут проверяет публичную папку Облака Mail.ru, скачивает новые .xlsx,
разбирает их и раздаёт расписание по схеме: /start -> дата -> класс -> буква.
Файлы без даты в названии, где внутри листа найдено несколько дней недели
(ПОНЕДЕЛЬНИК/ВТОРНИК/...), показываются отдельной кнопкой «Общее расписание»
и ведут в выбор дня недели, а дальше — как обычно, класс -> буква.
Кароче сами разберетесь с кодом если кому нужен ваще мой проэкт по факту там поменять только ссылку на ваш публичный мейл типо
 "https://cloud.mail.ru/public/"ключ ващего публичного облока"

ВНИМАНИЕ: этот файл сам по себе НЕ умеет ходить в Telegram через MTProto-
прокси. Если Telegram у вас доступен только через локальный MTProxy —
запускате schedule_bot_mtproto.py, а не этот файл напрямую (этот файл тогда
годится только для проверки парсера через --debug).

Установка:   pip install requests openpyxl
Запуск:      BOT_TOKEN=123:ABC python schedule_bot.py
Проверка парсера без бота:   python schedule_bot.py --debug
"""
import asyncio
import hashlib
import html as htmllib
import io
import importlib
import json
import logging
import os
import re
import sys
import threading
import time
from datetime import date, datetime, timedelta, timezone
from urllib.parse import quote

import openpyxl
import requests

def load_holidays():
    """Загружает календарь заново, чтобы изменения school_calendar.py
    подхватывались без старого значения, если модуль был уже импортирован."""
    try:
        import school_calendar as cal
        cal = importlib.reload(cal)
        return list(getattr(cal, "HOLIDAYS", []) or [])
    except (ImportError, AttributeError):
        return []


HOLIDAYS = load_holidays()

BOT_TOKEN = os.getenv("BOT_TOKEN", "")
FOLDER = "5U6H/ddj3h89S9"
PAGE_URL = f"https://cloud.mail.ru/public/{FOLDER}"
REFRESH_SECONDS = 600

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
    )
}

MONTHS = {
    "января": 1, "февраля": 2, "марта": 3, "апреля": 4, "мая": 5, "июня": 6,
    "июля": 7, "августа": 8, "сентября": 9, "октября": 10, "ноября": 11, "декабря": 12,
}
MONTHS_REV = {v: k for k, v in MONTHS.items()}

WEEKDAY_NAMES = ["понедельник", "вторник", "среда", "четверг", "пятница", "суббота", "воскресенье"]
WEEKDAY_ORDER = {name: i for i, name in enumerate(WEEKDAY_NAMES)}

# Екатеринбургское время: UTC+5. Так отсчёты не зависят от часового пояса Android/Termux.
EKB_TZ = timezone(timedelta(hours=5))


def today_ekb() -> date:
    return datetime.now(EKB_TZ).date()


def plural_days(n: int) -> str:
    n_abs = abs(n) % 100
    n1 = n_abs % 10
    if 11 <= n_abs <= 14:
        word = "дней"
    elif n1 == 1:
        word = "день"
    elif 2 <= n1 <= 4:
        word = "дня"
    else:
        word = "дней"
    return f"{n} {word}"


def holiday_status(today=None):
    """Текст для главного экрана: до ближайших каникул или до их конца."""
    today = today or today_ekb()
    parsed = []
    holidays = load_holidays()
    for title, start_s, end_s in holidays:
        try:
            start = date.fromisoformat(start_s)
            end = date.fromisoformat(end_s)
        except (TypeError, ValueError):
            continue
        if end < start:
            continue
        parsed.append((start, end, str(title)))
    parsed.sort()

    for start, end, title in parsed:
        if start <= today <= end:
            left = (end - today).days
            if left == 0:
                return f"🏖 {title}: последний день каникул"
            return f"🏖 {title}: до конца {plural_days(left)}"
        if today < start:
            left = (start - today).days
            return f"⏳ До каникул «{title}»: {plural_days(left)}"
    return None


def canon_weekday(word: str) -> str:
    return word[:1].upper() + word[1:]

# латинские буквы-двойники -> кириллица (в таблицах такое бывает)
LAT2CYR = str.maketrans("abcekmhopx", "авсекмнорх")
CLASS_RE = re.compile(r"^(\d{1,2})\s*([а-яёa-z]\d?)$")

log = logging.getLogger("schedule")

SESSION = requests.Session()
SESSION.trust_env = False   # не использовать системный прокси/VPN для Mail.ru


# ───────────────────────── 1. Облако Mail.ru ─────────────────────────

def list_folder():
    """Возвращает [(имя_файла, [список_возможных_ссылок_для_скачивания])]."""
    r = SESSION.get(PAGE_URL, headers=HEADERS, timeout=20)
    r.raise_for_status()
    page = r.text

    m = re.search(r'"weblink_get"\s*:\s*\[?\s*\{[^}]*?"url"\s*:\s*"([^"]+)"', page)
    if not m:
        raise RuntimeError("Не нашёл weblink_get на странице (Mail.ru поменял вёрстку?)")
    base = m.group(1).replace("\\/", "/").rstrip("/")

    files = {}
    dec = json.JSONDecoder()
    for lm in re.finditer(r'"list"\s*:\s*(?=\[)', page):
        try:
            arr, _ = dec.raw_decode(page, lm.end())
        except ValueError:
            continue
        for it in arr:
            if (
                isinstance(it, dict)
                and it.get("kind") == "file"
                and str(it.get("name", "")).lower().endswith(".xlsx")
            ):
                name = it["name"]
                links = []
                if it.get("weblink"):
                    links.append(f"{base}/{quote(it['weblink'])}")
                links.append(f"{base}/{quote(FOLDER + '/' + name)}")
                files[name] = links
    return list(files.items())


def download(links):
    last = None
    for url in links:
        try:
            r = SESSION.get(url, headers=HEADERS, timeout=30)
            if r.status_code == 200 and r.content[:2] == b"PK":  # xlsx = zip
                return r.content
            last = f"{r.status_code} {url}"
        except requests.RequestException as e:
            last = str(e)
    raise RuntimeError(f"Не удалось скачать файл: {last}")


# ───────────────────────── 2. Разбор xlsx ─────────────────────────

def to_int(v):
    try:
        return int(float(str(v).strip().replace(",", ".")))
    except (TypeError, ValueError):
        return None


def cval(ws, merges, r, c):
    """Значение ячейки с учётом объединения: для любой ячейки внутри
    объединённого диапазона возвращает значение его левой верхней ячейки."""
    key = (r, c)
    if key in merges:
        return merges[key]
    return ws.cell(r, c).value


def build_merge_map(ws):
    m = {}
    for rng in ws.merged_cells.ranges:
        top = ws.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                m[(r, c)] = top
    return m


# строки-баннеры без номера урока, которые не показываем как «событие»
# (обычные ежедневные надписи, неинтересные сами по себе)
IGNORE_BANNERS = {"разговоры о важном"}


def find_weekday(text: str):
    t = str(text or "").strip().lower()
    for name in WEEKDAY_NAMES:
        if t == name or t.startswith(name + " ") or t.startswith(name + "\n"):
            return name
    return None


def find_header_row(ws, r_start, r_end):
    limit = min(r_end, r_start + 20)
    for r in range(r_start, limit + 1):
        for c in range(1, ws.max_column + 1):
            if str(ws.cell(r, c).value or "").strip().lower() == "класс":
                return r
    return None


def parse_days(content: bytes):
    """-> {день_недели или None: {(класс, буква): {"shift","lessons","notices"}}}

    Один блок = одна табличка «класс/звонки/уроки» на листе. Обычный файл на
    один день даёт ровно один блок; файл-«мастер» на всю неделю (заголовки
    ПОНЕДЕЛЬНИК/ВТОРНИК/... внутри листа) — несколько."""
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    days_result: dict = {}

    for ws in wb.worksheets:
        merges = build_merge_map(ws)

        captions = []
        for r in range(1, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                wd = find_weekday(ws.cell(r, c).value)
                if wd:
                    captions.append((r, wd))
                    break

        if captions:
            blocks = []
            for i, (r, wd) in enumerate(captions):
                start = r + 1
                end = captions[i + 1][0] - 1 if i + 1 < len(captions) else ws.max_row
                blocks.append((wd, start, end))
        else:
            blocks = [(None, 1, ws.max_row)]

        for wd, r_start, r_end in blocks:
            header_row = find_header_row(ws, r_start, r_end)
            if not header_row:
                continue

            time_cols = [
                c for c in range(1, ws.max_column + 1)
                if str(ws.cell(header_row, c).value or "").strip().lower() == "класс"
            ]
            num_cols = [
                c for c in range(1, ws.max_column + 1)
                if "урока" in str(ws.cell(header_row, c).value or "").lower()
                or "урока" in str(ws.cell(header_row + 1, c).value or "").lower()
            ]

            for c in range(1, ws.max_column + 1):
                raw = str(ws.cell(header_row, c).value or "").strip().lower()
                m = CLASS_RE.match(raw)
                if not m:
                    continue
                grade = int(m.group(1))
                letter = m.group(2).translate(LAT2CYR)

                left = [t for t in time_cols if t < c]
                if not left:
                    continue
                tcol = max(left)
                nleft = [n for n in num_cols if n < tcol]
                ncol = max(nleft) if nleft else tcol - 1

                # события на несколько уроков подряд (вертикальное объединение
                # ячеек в колонке этого класса, например «экскурсия»)
                notices = []
                skip_rows = set()
                for rng in ws.merged_cells.ranges:
                    if (
                        rng.min_col <= c <= rng.max_col
                        and rng.max_row > rng.min_row
                        and rng.min_row > header_row
                        and rng.min_row >= r_start
                        and rng.max_row <= r_end
                    ):
                        text = str(ws.cell(rng.min_row, rng.min_col).value or "").strip()
                        if text:
                            notices.append(text)
                        skip_rows.update(range(rng.min_row, rng.max_row + 1))

                lessons = []
                for r in range(header_row + 1, r_end + 1):
                    if r in skip_rows:
                        continue
                    subj = str(cval(ws, merges, r, c) or "").strip()
                    if not subj:
                        continue
                    num = to_int(cval(ws, merges, r, ncol)) if ncol >= 1 else None
                    t = str(cval(ws, merges, r, tcol) or "").strip()
                    if num is None:
                        # строка без номера урока: обычную «Разговоры о важном»
                        # молчим, а прочие подписи (Медосмотр, «Россия — мои
                        # горизонты» и т.п.) показываем отдельным событием
                        if t and subj.lower() not in IGNORE_BANNERS:
                            notices.append(subj)
                        continue
                    lessons.append((num, t, subj))

                key = (grade, letter)
                day_dict = days_result.setdefault(wd, {})
                if key not in day_dict or lessons or notices:
                    day_dict[key] = {"shift": ws.title, "lessons": lessons, "notices": notices}

    return days_result


def date_from_name(name: str):
    m = re.search(r"(\d{1,2})\s+([а-яё]+)\s+(\d{4})", name.lower())
    if m and m.group(2) in MONTHS:
        try:
            return date(int(m.group(3)), MONTHS[m.group(2)], int(m.group(1)))
        except ValueError:
            pass
    return None


# ───────────────────────── форматирование ответа ─────────────────────────

# расшифровка сокращений (дополняй сам)
ABBR = {
    "русс": "русский", "русск": "русский", "литер": "литература", "лит": "литература",
    "матем": "математика", "геогр": "география", "биолог": "биология",
    "истор": "история", "инф": "информатика", "информ": "информатика",
    "ин": "иностранный язык", "англ": "английский", "англ.язык": "английский",
    "нем": "немецкий", "фк": "физкультура", "физ-ра": "физкультура",
    "ит": "ИТ",
}


def fmt_time(t: str) -> str:
    ts = re.findall(r"(\d{1,2})[-:.](\d{2})", t)
    if len(ts) == 2:
        return f"{ts[0][0]}:{ts[0][1]} - {ts[1][0]}:{ts[1][1]}"
    return t


def fmt_subject(s: str) -> str:
    m = re.match(r"^(/?[^\d/]+?)\s*(\d+)$", s.strip())   # «русс 24» -> «русский 24 каб»
    if m:
        name = m.group(1).strip()
        key = name.lstrip("/").lower()
        name = ABBR.get(key, name)
        return f"{name} {m.group(2)} каб"
    return s


# ───────────────────────── 3. Кэш расписаний ─────────────────────────

# имя файла -> {"sid", "title", "sort", "data"} — обычные файлы на один день
SCHEDULES: dict = {}
# имя файла -> {"key", "title", "days": {день_недели: {...}}} — файлы-мастера на неделю
WEEK_FILES: dict = {}


def sync_refresh():
    files = list_folder()
    names_now = {n for n, _ in files}
    for name, links in files:
        if name in SCHEDULES or name in WEEK_FILES:
            continue
        log.info("Новый файл: %s", name)
        days = parse_days(download(links))
        d = date_from_name(name)
        if d is None and len(days) > 1:
            # файл-мастер сразу на несколько дней недели («ПРОЕКТ» и т.п.)
            WEEK_FILES[name] = {
                "key": hashlib.md5(name.encode()).hexdigest()[:8],
                "title": name.rsplit(".", 1)[0],
                "days": days,
            }
        else:
            data = next(iter(days.values()), {})
            SCHEDULES[name] = {
                "sid": d.isoformat() if d else "f" + hashlib.md5(name.encode()).hexdigest()[:8],
                "title": f"{d.day} {MONTHS_REV[d.month]} {d.year}" if d else name.rsplit(".", 1)[0],
                "sort": d.isoformat() if d else name,
                "data": data,
            }
    for name in list(SCHEDULES):        # удалённые из папки — убираем и у себя
        if name not in names_now:
            log.info("Файл удалён из папки: %s", name)
            del SCHEDULES[name]
    for name in list(WEEK_FILES):
        if name not in names_now:
            log.info("Файл удалён из папки: %s", name)
            del WEEK_FILES[name]


def week_file_by_key(key):
    for wf in WEEK_FILES.values():
        if wf["key"] == key:
            return wf
    return None


def by_sid(sid):
    if sid.startswith("w~"):
        _, file_key, wd = sid.split("~", 2)
        wf = week_file_by_key(file_key)
        if not wf or wd not in wf["days"]:
            return None
        return {"sid": sid, "title": f"{canon_weekday(wd)} · {wf['title']}", "data": wf["days"][wd]}
    for s in SCHEDULES.values():
        if s["sid"] == sid:
            return s
    return None


def home_target(sid):
    """Куда ведёт кнопка «домой»: к списку дат или к списку дней недели файла-мастера."""
    if sid.startswith("w~"):
        file_key = sid.split("~", 2)[1]
        return f"wk:{file_key}"
    return "back"


# ───────────────────────── 3.5. Поиск ближайшего урока ─────────────────────────

def primary_week_file():
    """Выбирает самый похожий на главный недельный файл: больше дней/классов."""
    if not WEEK_FILES:
        return None
    return max(
        WEEK_FILES.values(),
        key=lambda wf: (
            len(wf.get("days", {})),
            sum(len(day) for day in wf.get("days", {}).values()),
            wf.get("title", ""),
        ),
    )


def find_classes_data():
    wf = primary_week_file()
    return wf, (wf.get("days", {}) if wf else {})




def lesson_start_minutes(t: str):
    """Начало урока в минутах от полуночи; None, если время не распознано."""
    m = re.search(r"(\d{1,2})[-:.](\d{2})", str(t or ""))
    if not m:
        return None
    return int(m.group(1)) * 60 + int(m.group(2))


def subject_rooms(raw):
    """Достаёт номера кабинетов из исходной записи предмета."""
    nums = re.findall(r"\d+", str(raw or ""))
    out = []
    for n in nums:
        if n not in out:
            out.append(n)
    return out


def _subject_parts(raw):
    """Разбирает ячейку на части по '/'.

    Важное правило для школьной таблицы:
      левая часть до '/' = 1-я подгруппа,
      правая часть после '/' = 2-я подгруппа.

    Поэтому 'ин(м)/ит' — это два разных предмета: иностранный язык у 1-й
    подгруппы и ИТ у 2-й. Они должны быть двумя отдельными кнопками.
    """
    return str(raw or '').strip().split('/')


def _clean_subject_part(part):
    """Оставляет только название предмета из одной части ячейки."""
    part = str(part or '').strip()
    if not part:
        return ''
    part = re.sub(r'\d+', ' ', part)
    part = re.sub(r'\bакт(?:овый)?\.?\s*зал\b', ' ', part, flags=re.I)
    part = re.sub(r'\b(?:каб(?:инет)?|зал)\b', ' ', part, flags=re.I)
    # (м), (е), (с) и т.п. — инициалы/метка учителя.
    part = re.sub(r'\s*\([^)]+\)\s*', ' ', part)
    part = re.sub(r'\s+', ' ', part).strip(' /-—_\t')
    if not part:
        return ''
    key = part.lstrip('/').casefold()
    return ABBR.get(key, part.lstrip('/'))


def subject_label(raw):
    """Возвращает название одного предмета. Для ячейки с '/'
    используется первая непустая часть; список предметов строится отдельно
    по каждой части."""
    for part in _subject_parts(raw):
        label = _clean_subject_part(part)
        if label:
            return label
    return ''


def subject_parts_with_groups(raw):
    """[(label, subgroup, raw_part, index)] для всех непустых частей."""
    parts = _subject_parts(raw)
    out = []
    for idx, part in enumerate(parts):
        label = _clean_subject_part(part)
        if not label:
            continue
        subgroup = None if '/' not in str(raw or '') else (1 if idx == 0 else 2)
        out.append((label, subgroup, part, idx))
    return out


def subject_subgroups(raw):
    """Подгруппа предмета в исходной ячейке: (), (1,), (2,) или (1, 2)."""
    groups = [g for _label, g, _part, _idx in subject_parts_with_groups(raw) if g is not None]
    if not groups:
        return ()
    return tuple(sorted(set(groups)))


def subject_display(raw):
    label = subject_label(raw)
    groups = subject_subgroups(raw)
    if not groups:
        return label
    if groups == (1,):
        return f'{label} — 1 подгруппа'
    if groups == (2,):
        return f'{label} — 2 подгруппа'
    return f'{label} — обе подгруппы'


def _part_rooms(part):
    nums = re.findall(r'\d+', str(part or ''))
    out = []
    for n in nums:
        if n not in out:
            out.append(n)
    return out


def subjects_for_class(grade, letter):
    """Уникальные предметы класса. Каждая сторона '/' — отдельный предмет."""
    _, days = find_classes_data()
    found = {}
    for wd in WEEKDAY_NAMES:
        info = days.get(wd, {}).get((int(grade), letter))
        if not info:
            continue
        for _num, _t, subj in info.get('lessons', []):
            for label, _group, _part, _idx in subject_parts_with_groups(subj):
                found.setdefault(label.casefold(), label)
    return [found[k] for k in sorted(found)]


def subject_match_parts(raw, target):
    """Все части ячейки, соответствующие выбранному предмету (может быть
    несколько — например, один и тот же предмет одновременно у обеих
    подгрупп, просто в разных кабинетах)."""
    target_label = str(target or '').strip().casefold()
    return [
        (label, group, part, idx)
        for label, group, part, idx in subject_parts_with_groups(raw)
        if label.casefold() == target_label
    ]


def subject_matches(raw, target):
    return bool(subject_match_parts(raw, target))



def screen_find_grades():
    _, days = find_classes_data()
    grades = sorted({g for day in days.values() for g, _ in day})
    if not grades:
        return "🔎 В общем недельном файле пока нет классов.", kb([], extra_rows=[[btn("🏠 К датам", "back")]])
    return (
        "🔎 <b>Какой предмет ближайший?</b>\nВыбери класс:\n\n<i>Поиск ориентируется по общему файлу со всем расписанием. Рекомендуем перепроверять результат по расписанию на конкретную дату.</i>",
        kb([btn(str(g), f"fg:{g}") for g in grades], extra_rows=[[btn("⬅️ К датам", "back")]]),
    )


def screen_find_letters(grade):
    _, days = find_classes_data()
    letters = sorted({l for day in days.values() for g, l in day if g == int(grade)})
    if not letters:
        return None
    return (
        f"🔎 Класс {grade}. Выбери букву:",
        kb([btn(f"{grade}{l}", f"fl:{grade}:{l}") for l in letters], extra_rows=[[btn("⬅️ Назад", "find")]]),
    )


def screen_find_subjects(grade, letter):
    subjects = subjects_for_class(grade, letter)
    if not subjects:
        return (
            f"🔎 Для {grade}{htmllib.escape(letter)} предметы не нашлись.",
            kb([], extra_rows=[[btn("⬅️ Назад", f"fg:{grade}"), btn("🏠 К датам", "back")]]),
        )
    buttons = [btn(subj, f"fs:{grade}:{letter}:{i}") for i, subj in enumerate(subjects)]
    return (
        f"🔎 <b>{grade}{htmllib.escape(letter)}</b>\nКакой предмет ищем?\n\n<i>Бот ориентируется по общему файлу со всем расписанием. Рекомендуем перепроверять результат по расписанию на конкретную дату.</i>",
        kb(buttons, per_row=2, extra_rows=[[btn("⬅️ Назад", f"fg:{grade}"), btn("🏠 К датам", "back")]]),
    )


def screen_next_lesson(grade, letter, subject_index):
    wf, days = find_classes_data()
    subjects = subjects_for_class(grade, letter)
    try:
        target = subjects[int(subject_index)]
    except (ValueError, IndexError):
        return None

    now = datetime.now(EKB_TZ)
    today = now.date()
    today_idx = today.weekday()
    now_minutes = now.hour * 60 + now.minute
    hits = []

    for wd, day_data in days.items():
        info = day_data.get((int(grade), letter))
        if not info:
            continue
        wd_idx = WEEKDAY_ORDER.get(wd)
        if wd_idx is None:
            continue
        offset = (wd_idx - today_idx) % 7
        lesson_date = today + timedelta(days=offset)

        for num, t, subj in info.get('lessons', []):
            matches = subject_match_parts(subj, target)
            if not matches:
                continue
            start_min = lesson_start_minutes(t)
            if offset == 0 and start_min is not None and start_min <= now_minutes:
                continue
            hits.append((offset, lesson_date, wd, num, t, matches, start_min or 9999))

    if not hits:
        body = (
            f'Не нашёл <b>{htmllib.escape(str(target))}</b> у '
            f'{grade}{htmllib.escape(letter)} в недельном файле.'
        )
    else:
        offset, lesson_date, wd, num, t, matches, _start_min = min(
            hits, key=lambda x: (x[0], x[6], x[3])
        )
        when = 'сегодня' if offset == 0 else ('завтра' if offset == 1 else f'через {plural_days(offset)}')
        label = matches[0][0]

        if len(matches) == 1:
            _label, group, raw_part, _idx = matches[0]
            groups_line = ''
            if group == 1:
                groups_line = '\n👥 1 подгруппа'
            elif group == 2:
                groups_line = '\n👥 2 подгруппа'
            rooms = _part_rooms(raw_part)
            room_line = f"\n🚪 Кабинет: {htmllib.escape(' / '.join(rooms))}" if rooms else ''
        else:
            # предмет одновременно у нескольких частей (обычно — у обеих
            # подгрупп сразу, просто в разных кабинетах)
            sub_lines = []
            for _label, group, raw_part, _idx in sorted(matches, key=lambda m: (m[1] is None, m[1])):
                rooms = _part_rooms(raw_part)
                room_txt = htmllib.escape(' / '.join(rooms)) if rooms else '—'
                gtxt = f'{group} подгруппа' if group else 'без подгруппы'
                sub_lines.append(f'👥 {gtxt} — каб. {room_txt}')
            groups_line = '\n' + '\n'.join(sub_lines)
            room_line = ''

        body = (
            f'🔎 <b>{grade}{htmllib.escape(letter)} · {htmllib.escape(label)}</b>\n'
            f'Ближайший урок: <b>{canon_weekday(wd)}, {lesson_date.day} {MONTHS_REV[lesson_date.month]}</b> ({when})\n'
            f'⏰ Урок №{num} · {htmllib.escape(fmt_time(t))}{groups_line}{room_line}'
        )
        if wf:
            body += f"\n\n<i>По файлу: {htmllib.escape(wf['title'])}</i>"
    return body, kb([], extra_rows=[[btn('⬅️ К предметам', f'fl:{grade}:{letter}'), btn('🏠 К датам', 'back')]])


def screen_holidays():
    holidays = load_holidays()
    parsed = []
    for title, start_s, end_s in holidays:
        try:
            start = date.fromisoformat(start_s)
            end = date.fromisoformat(end_s)
        except (TypeError, ValueError):
            continue
        if end >= start:
            parsed.append((start, end, str(title)))
    parsed.sort()
    if not parsed:
        return (
            '🏖 <b>Каникулы</b>\n\n'
            'Даты каникул пока не заполнены в <code>school_calendar.py</code>.\n'
            'После изменения файла перезапусти бота — календарь перечитается.',
            kb([], extra_rows=[[btn('⬅️ К датам', 'back')]])
        )
    today = today_ekb()
    lines = ['🏖 <b>Каникулы</b>']
    current_or_next = None
    for start, end, title in parsed:
        if start <= today <= end:
            left = (end - today).days
            current_or_next = f'🏖 Сейчас: <b>{htmllib.escape(title)}</b> — до конца {plural_days(left)}'
            break
        if today < start and current_or_next is None:
            current_or_next = f'⏳ До «{htmllib.escape(title)}»: {plural_days((start - today).days)}'
            break
    if current_or_next:
        lines.append('\n' + current_or_next)
    lines.append('')
    for start, end, title in parsed:
        lines.append(f'• <b>{htmllib.escape(title)}</b>: {start.day:02d}.{start.month:02d}.{start.year} — {end.day:02d}.{end.month:02d}.{end.year}')
    return '\n'.join(lines), kb([], extra_rows=[[btn('⬅️ К датам', 'back')]])


# ───────────────────────── 4. Бот (без aiogram, только requests) ─────────────────────────

TG = requests.Session()   # для Telegram системный VPN/прокси разрешён
if os.getenv("TG_PROXY"):     # например socks5h://127.0.0.1:1080 (нужен pip install pysocks)
    TG.proxies = {"http": os.environ["TG_PROXY"], "https": os.environ["TG_PROXY"]}


def tg(method, **params):
    r = TG.post(f"https://api.telegram.org/bot{BOT_TOKEN}/{method}", json=params, timeout=70)
    data = r.json()
    if not data.get("ok"):
        raise RuntimeError(f"{method}: {data.get('description')}")
    return data["result"]


def btn(text, data):
    return {"text": text, "callback_data": data}


def kb(buttons, per_row=4, extra_rows=()):
    rows = [buttons[i:i + per_row] for i in range(0, len(buttons), per_row)]
    rows.extend(extra_rows)
    return {"inline_keyboard": rows}


def screen_dates():
    items = sorted(list(SCHEDULES.values()), key=lambda x: x["sort"])
    buttons = [btn(s["title"], f"d:{s['sid']}") for s in items]
    extra = [[btn("🔎 Когда следующий урок?", "find")]]
    extra += [[btn("🏖 Каникулы", "hol")]]
    extra += [[btn(f"📚 {wf['title']}", f"wk:{wf['key']}")] for wf in WEEK_FILES.values()]
    if not buttons and len(extra) == 1:
        return "Пока расписаний в папке нет — попробуй чуть позже.", None
    markup = kb(buttons, per_row=2, extra_rows=extra)
    status = holiday_status()
    head = "За какое число нужно расписание?"
    if status:
        head = f"{status}\n\n{head}"
    return head, markup


def screen_weekdays(file_key):
    wf = week_file_by_key(file_key)
    if not wf:
        return None
    days_sorted = sorted(wf["days"].keys(), key=lambda w: WEEKDAY_ORDER.get(w, 99))
    buttons = [btn(canon_weekday(w), f"d:w~{file_key}~{w}") for w in days_sorted]
    markup = kb(buttons, per_row=2, extra_rows=[[btn("⬅️ К датам", "back")]])
    return f"📚 {htmllib.escape(wf['title'])}\nВыбери день недели:", markup


def screen_grades(sid):
    s = by_sid(sid)
    if not s:
        return None
    grades = sorted({g for g, _ in s["data"]})
    markup = kb(
        [btn(str(g), f"g:{sid}:{g}") for g in grades],
        extra_rows=[[btn("⬅️ К датам", home_target(sid))]],
    )
    return f"📅 {htmllib.escape(s['title'])}\nВыбери класс:", markup


def screen_letters(sid, grade):
    s = by_sid(sid)
    if not s:
        return None
    letters = sorted({l for g, l in s["data"] if g == int(grade)})
    markup = kb(
        [btn(f"{grade}{l}", f"c:{sid}:{grade}:{l}") for l in letters],
        extra_rows=[[btn("⬅️ Назад", f"d:{sid}")]],
    )
    return f"📅 {htmllib.escape(s['title'])}\nКласс {grade}. Выбери букву:", markup


def screen_schedule(sid, grade, letter):
    s = by_sid(sid)
    if not s:
        return None
    info = s["data"].get((int(grade), letter))
    lines = [f"<b>{grade}{htmllib.escape(letter)}</b>  ({htmllib.escape(s['title'])})"]
    if info:
        for note in info.get("notices", []):
            lines.append(f"📌 {htmllib.escape(note)}")
        if info["lessons"]:
            for num, t, subj in info["lessons"]:
                lines.append(
                    f"{num}. {htmllib.escape(fmt_time(t))} {htmllib.escape(fmt_subject(subj))}"
                )
        elif not info.get("notices"):
            lines.append("Уроков в этот день нет 🎉")
    else:
        lines.append("Для этого класса данных нет.")
    markup = kb(
        [],
        extra_rows=[[btn("⬅️ Другой класс", f"d:{sid}"), btn("🏠 К датам", home_target(sid))]],
    )
    return "\n".join(lines), markup


def handle_message(msg):
    text = (msg.get("text") or "").strip()
    chat_id = msg["chat"]["id"]
    if text.startswith("/start"):
        body, markup = screen_dates()
        if markup:
            body = "Привет! 👋 Я показываю расписание уроков.\n\n" + body
        tg("sendMessage", chat_id=chat_id, text=body, reply_markup=markup, parse_mode="HTML")
    elif text.startswith("/update"):
        sync_refresh()
        tg("sendMessage", chat_id=chat_id, text=f"Обновлено. Расписаний в базе: {len(SCHEDULES)}")


def handle_callback(cb):
    kind, *a = cb["data"].split(":")
    if kind == "back":
        res = screen_dates()
    elif kind == "find":
        res = screen_find_grades()
    elif kind == "hol":
        res = screen_holidays()
    elif kind == "fg":
        res = screen_find_letters(a[0])
    elif kind == "fl":
        res = screen_find_subjects(a[0], a[1])
    elif kind == "fs":
        res = screen_next_lesson(a[0], a[1], a[2])
    elif kind == "wk":
        res = screen_weekdays(a[0])
    elif kind == "d":
        res = screen_grades(a[0])
    elif kind == "g":
        res = screen_letters(a[0], a[1])
    elif kind == "c":
        res = screen_schedule(a[0], a[1], a[2])
    else:
        res = None
    if res is None:
        tg("answerCallbackQuery", callback_query_id=cb["id"],
           text="Это расписание уже удалили, нажми /start", show_alert=True)
        return
    tg("answerCallbackQuery", callback_query_id=cb["id"])
    text, markup = res
    m = cb["message"]
    try:
        tg("editMessageText", chat_id=m["chat"]["id"], message_id=m["message_id"],
           text=text, reply_markup=markup or {"inline_keyboard": []}, parse_mode="HTML")
    except RuntimeError as e:
        if "not modified" not in str(e):
            raise


def refresher():
    while True:
        try:
            sync_refresh()
        except Exception:
            log.exception("Не удалось обновить папку")
        time.sleep(REFRESH_SECONDS)


def run_bot():
    threading.Thread(target=refresher, daemon=True).start()
    while True:
        try:
            me = tg("getMe")
            break
        except requests.RequestException as e:
            log.error("Нет связи с Telegram (включи VPN на устройстве?): %s", e.__class__.__name__)
            time.sleep(10)
        except RuntimeError as e:
            sys.exit(f"Telegram отклонил токен: {e}")
    log.info("Бот запущен: @%s", me.get("username"))
    offset = None
    while True:
        params = {"timeout": 50, "allowed_updates": ["message", "callback_query"]}
        if offset:
            params["offset"] = offset
        try:
            updates = tg("getUpdates", **params)
        except Exception:
            log.exception("Ошибка связи с Telegram, пробую снова")
            time.sleep(3)
            continue
        for u in updates:
            offset = u["update_id"] + 1
            try:
                if "callback_query" in u:
                    handle_callback(u["callback_query"])
                elif "message" in u:
                    handle_message(u["message"])
            except Exception:
                log.exception("Ошибка обработки обновления")


# ───────────────────────── запуск ─────────────────────────

def debug():
    for n, links in list_folder():
        days = parse_days(download(links))
        d = date_from_name(n)
        kind = "файл-мастер на неделю" if (d is None and len(days) > 1) else "обычный день"
        print(f"\n=== {n} ({kind}), дней найдено: {len(days)}")
        for wd, data in sorted(days.items(), key=lambda kv: WEEKDAY_ORDER.get(kv[0], -1)):
            print(f"--- {canon_weekday(wd) if wd else '(без подписи дня)'}: найдено классов {len(data)}")
            by_shift = {}
            for (g, l), v in sorted(data.items()):
                by_shift.setdefault(v["shift"], []).append(f"{g}{l}")
            for shift, cl in by_shift.items():
                print(f"[{shift}] {len(cl)}: {' '.join(cl)}")
            with_notices = {f"{g}{l}": v["notices"] for (g, l), v in data.items() if v["notices"]}
            if with_notices:
                print("Заметки (без номера урока):", with_notices)
            empty = [f"{g}{l}" for (g, l), v in sorted(data.items()) if not v["lessons"] and not v["notices"]]
            print("Классы без уроков:", " ".join(empty) or "нет")
            for k in [(7, "г"), (8, "ю")]:
                if k in data:
                    print(f"{k[0]}{k[1]}:", data[k]["lessons"], data[k]["notices"])


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO)
    if "--debug" in sys.argv:
        debug()
    else:
        if not BOT_TOKEN:
            sys.exit("Задай токен: BOT_TOKEN=... python schedule_bot.py")
        run_bot()
